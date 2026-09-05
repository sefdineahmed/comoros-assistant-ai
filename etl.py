# -*- coding: utf-8 -*-
"""
etl.py — Jalon 1 de l'atelier, version production.

EXTRACT  : lire chaque fichier (PDF/DOCX/PPTX/XLSX) et en tirer du texte brut.
TRANSFORM: nettoyer + découper ce texte en passages ("chunks") courts.
LOAD     : sauvegarder le résultat sur disque (cache), pour que les étapes
           suivantes (entraînement, tests, app) n'aient jamais à relire
           les fichiers sources.

Exécuter directement ce fichier lance l'ETL complet :
    python etl.py
"""

import os
import pickle
import hashlib

from tqdm import tqdm

import config


def nettoyer_texte(texte):
    """Supprime les caractères Unicode invalides ('surrogates' orphelins)
    que produisent certains vieux PDF à l'encodage de police propriétaire
    (ex: SymbolSetEncoding). Sans ce nettoyage, ces caractères passent
    l'extraction sans erreur mais font planter tout ce qui essaie
    d'encoder le texte en UTF-8 plus tard (hash, sauvegarde, etc.)."""
    if not texte:
        return ""
    return texte.encode("utf-8", errors="ignore").decode("utf-8")


# =============================================================================
# EXTRACT — un extracteur de texte par type de fichier
# =============================================================================

def extraire_pdf(chemin):
    from pypdf import PdfReader
    try:
        pages = PdfReader(chemin).pages
        texte = "\n".join((p.extract_text() or "") for p in pages)
        return nettoyer_texte(texte)
    except Exception as e:
        print(f"   ⚠️ PDF illisible ({e}) — probablement un scan sans OCR : {chemin}")
        return ""


def extraire_docx(chemin):
    from docx import Document
    try:
        doc = Document(chemin)
        texte = "\n".join(p.text for p in doc.paragraphs)
        return nettoyer_texte(texte)
    except Exception as e:
        print(f"   ⚠️ DOCX illisible ({e}) : {chemin}")
        return ""


def extraire_pptx(chemin):
    from pptx import Presentation
    try:
        prs = Presentation(chemin)
        textes = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    textes.append(shape.text)
        return nettoyer_texte("\n".join(textes))
    except Exception as e:
        print(f"   ⚠️ PPTX illisible ({e}) : {chemin}")
        return ""


def extraire_xlsx(chemin):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
        textes = []
        for feuille in wb.worksheets:
            for ligne in feuille.iter_rows(values_only=True):
                cellules = [str(c) for c in ligne if c is not None]
                if cellules:
                    textes.append(" | ".join(cellules))
        return nettoyer_texte("\n".join(textes))
    except Exception as e:
        print(f"   ⚠️ XLSX illisible ({e}) : {chemin}")
        return ""


EXTRACTEURS = {
    ".pdf": extraire_pdf,
    ".docx": extraire_docx,
    ".pptx": extraire_pptx,
    ".xlsx": extraire_xlsx,
    ".xls": extraire_xlsx,
}


def lister_fichiers(dossier_racine):
    """Parcourt récursivement le dossier. La 'catégorie' d'un fichier est le
    premier sous-dossier sous la racine (ex: 'Santé', 'Pêche', 'Religion')."""
    fichiers = []
    for racine, _, noms in os.walk(dossier_racine):
        for nom in noms:
            ext = os.path.splitext(nom)[1].lower()
            if ext not in config.EXTENSIONS_SUPPORTEES:
                continue
            chemin = os.path.join(racine, nom)
            relatif = os.path.relpath(chemin, dossier_racine)
            categorie = relatif.split(os.sep)[0] if os.sep in relatif else "Racine"
            fichiers.append((chemin, categorie, nom))
    return fichiers


def hash_fichier(chemin):
    """Empreinte rapide (taille + date de modif) : permet de ne réextraire
    que les fichiers nouveaux ou modifiés depuis la dernière exécution."""
    stat = os.stat(chemin)
    return hashlib.md5(f"{chemin}-{stat.st_size}-{stat.st_mtime}".encode()).hexdigest()


def extract(dossier_racine=None):
    """EXTRACT : lit tous les fichiers supportés et renvoie
    {chemin: {texte, categorie, nom, chemin}}. Incrémental via cache."""
    dossier_racine = dossier_racine or config.DOSSIER_DOCUMENTS

    cache = {}
    if os.path.exists(config.CACHE_EXTRACTION):
        with open(config.CACHE_EXTRACTION, "rb") as f:
            cache = pickle.load(f)
        print(f"📦 Cache d'extraction trouvé : {len(cache)} documents déjà extraits.")

    chemin_absolu = os.path.abspath(dossier_racine)

    if not os.path.exists(dossier_racine):
        raise FileNotFoundError(
            f"Le dossier n'existe pas : '{dossier_racine}'\n"
            f"   (chemin absolu résolu : {chemin_absolu})\n"
            "   Vérifie config.DOSSIER_DOCUMENTS — le chemin est relatif au "
            "dossier depuis lequel tu lances 'python main.py' / 'python etl.py'."
        )

    fichiers = lister_fichiers(dossier_racine)
    print(f"🔍 {len(fichiers)} fichiers trouvés dans {chemin_absolu}")

    if not fichiers:
        # Le dossier existe mais rien de supporté dedans : on aide au diagnostic
        # en listant ce qu'il contient réellement (fichiers + sous-dossiers).
        contenu = os.listdir(dossier_racine)
        raise FileNotFoundError(
            f"Le dossier '{chemin_absolu}' existe mais ne contient aucun fichier "
            f"avec une extension supportée {sorted(config.EXTENSIONS_SUPPORTEES)}.\n"
            f"   Contenu trouvé à la racine ({len(contenu)} élément(s)) : {contenu[:10]}"
            + (" ..." if len(contenu) > 10 else "")
        )

    resultats = {}
    nouveaux = 0
    for chemin, categorie, nom in tqdm(fichiers, desc="Extraction"):
        cle = hash_fichier(chemin)
        if cle in cache:
            resultats[chemin] = cache[cle]
            continue
        ext = os.path.splitext(nom)[1].lower()
        texte = EXTRACTEURS[ext](chemin)
        entree = {"texte": texte, "categorie": categorie, "nom": nom, "chemin": chemin}
        resultats[chemin] = entree
        cache[cle] = entree
        nouveaux += 1

    with open(config.CACHE_EXTRACTION, "wb") as f:
        pickle.dump(cache, f)

    print(f"✅ EXTRACT terminé : {len(resultats)} documents disponibles ({nouveaux} nouveaux).")
    vides = [r["nom"] for r in resultats.values() if len(r["texte"].strip()) < config.LONGUEUR_MIN_DOCUMENT]
    if vides:
        print(f"⚠️ {len(vides)} document(s) quasi vides (scan sans OCR probable), ex: {vides[:5]}")
    return resultats


# =============================================================================
# TRANSFORM — nettoyage + découpage en passages
# =============================================================================

def decouper(texte, taille=None, chevauchement=None):
    """Découpe un texte en passages, en coupant de préférence en fin de phrase."""
    taille = taille or config.TAILLE_CHUNK
    chevauchement = chevauchement or config.CHEVAUCHEMENT

    texte = " ".join(texte.split())
    passages, debut = [], 0
    while debut < len(texte):
        fin = min(debut + taille, len(texte))
        if fin < len(texte):
            coupe = texte.rfind(". ", debut + taille // 2, fin)
            if coupe != -1:
                fin = coupe + 1
        passages.append(texte[debut:fin].strip())
        if fin >= len(texte):
            break
        debut = max(fin - chevauchement, debut + 1)
    return [p for p in passages if p]


def transform(docs_bruts):
    """TRANSFORM : {chemin: {texte, categorie, nom}} → liste de passages,
    chacun gardant sa source complète pour la citation."""
    DOCUMENTS = []
    for chemin, info in docs_bruts.items():
        # Nettoyage défensif : même si le texte vient du cache (extrait avant
        # ce correctif) et contient encore des caractères invalides, on les
        # retire ici avant le découpage.
        texte_propre = nettoyer_texte(info["texte"])
        if len(texte_propre.strip()) < config.LONGUEUR_MIN_DOCUMENT:
            continue
        titre_doc = os.path.splitext(info["nom"])[0]
        for i, p in enumerate(decouper(texte_propre), 1):
            DOCUMENTS.append({
                "titre": f"{titre_doc} · passage {i}",
                "texte": p,
                "categorie": info["categorie"],
                "fichier": info["nom"],
                "chemin": info["chemin"],
            })
    print(f"✅ TRANSFORM terminé : {len(DOCUMENTS)} passages construits.")
    return DOCUMENTS


# =============================================================================
# LOAD — persistance sur disque
# =============================================================================

def load(DOCUMENTS):
    """LOAD : sauvegarde les passages sur disque pour les étapes suivantes."""
    with open(config.CACHE_PASSAGES, "wb") as f:
        pickle.dump(DOCUMENTS, f)
    print(f"✅ LOAD terminé : passages sauvegardés dans {config.CACHE_PASSAGES}")


def charger_passages():
    """Utilisé par les autres modules (entrainement.py, app.py, tests) pour
    récupérer les passages sans relancer tout l'ETL."""
    if not os.path.exists(config.CACHE_PASSAGES):
        raise FileNotFoundError(
            "Aucun passage trouvé. Lance d'abord l'ETL : python etl.py"
        )
    with open(config.CACHE_PASSAGES, "rb") as f:
        return pickle.load(f)


def run_etl(dossier_racine=None):
    """Pipeline complet Extract → Transform → Load."""
    docs_bruts = extract(dossier_racine)
    DOCUMENTS = transform(docs_bruts)
    load(DOCUMENTS)
    return DOCUMENTS


if __name__ == "__main__":
    run_etl()
